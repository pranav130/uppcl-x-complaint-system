"""Builds an exact-match replica of the original UPPCL Google Sheet.

Matches:
- Headers in Row 1 (Height: 66px, Background: #4DD0E1 Cyan, Font: Verdana 13pt Bold, Color: #1A2F40)
- Counter row in Row 2 (Height: 33px, Background: #4DD0E1)
- Frozen Panes at Row 3 (Freeze Rows 1 & 2)
- Data starting at Row 3 with solid black borders and Verdana typography
- Pre-populated with representative trial complaints
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from uppcl_x_automation.config import load_config
from uppcl_x_automation.sheets_client import EXISTING_HEADERS, AUTOMATION_HEADERS, FULL_HEADERS


def build_exact_format_workbook(output_path: str = "UPPCL_X_Complaint_Trial_Workbook.xlsx") -> None:
    wb = openpyxl.Workbook()
    config = load_config()

    # Exact styles from original sheet
    cyan_fill = PatternFill(start_color="4DD0E1", end_color="4DD0E1", fill_type="solid")
    automation_fill = PatternFill(start_color="00ACC1", end_color="00ACC1", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    header_font = Font(name="Verdana", size=13, bold=True, color="1A2F40")
    data_font = Font(name="Verdana", size=11, color="1A2F40")
    link_font = Font(name="Times New Roman", size=11, color="1D9BF0", underline="single")
    bold_data_font = Font(name="Verdana", size=11, bold=True, color="1A2F40")

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    black_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    # -------------------------------------------------------------------------
    # TAB 1: Twitter Complaint List
    # -------------------------------------------------------------------------
    ws_main = wb.active
    ws_main.title = "Twitter Complaint List"

    # ROW 1: Headers
    ws_main.row_dimensions[1].height = 66
    for col_idx, h_text in enumerate(FULL_HEADERS, start=1):
        cell = ws_main.cell(row=1, column=col_idx, value=h_text)
        cell.font = header_font
        cell.alignment = center_align
        cell.fill = automation_fill if col_idx > len(EXISTING_HEADERS) else cyan_fill
        cell.border = black_border

    # ROW 2: Counter / Summary Row (as in original sheet)
    ws_main.row_dimensions[2].height = 33
    for col_idx in range(1, len(FULL_HEADERS) + 1):
        cell = ws_main.cell(row=2, column=col_idx)
        cell.font = bold_data_font
        cell.alignment = center_align
        cell.fill = cyan_fill
        cell.border = black_border
    
    # Put sample count in column J (Complaint Type counter)
    ws_main.cell(row=2, column=10, value="=COUNTA(E3:E50)")
    # Column Q (REPLY_WITHIN_60_MIN counter)
    ws_main.cell(row=2, column=17, value="=COUNTIF(Q3:Q50, \"YES\")")

    # Freeze panes at Row 3 (Freeze Rows 1 & 2)
    ws_main.freeze_panes = "A3"

    # ROW 3 Onwards: Trial Data Rows
    sample_rows = [
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/Diwaker70180261/status/2111604765486707028?s=20",
            "Diwaker70180261", "upm", "", "YES", "SUPPLY_OUTAGE",
            "MVVNL", "DEVIPATAN", "GONDA", "Verified match for GONDA -> MVVNL",
            "Forwarded to Discom", "YES", "", "", "",
            "2111604765486707028",
            "@UPPCLLKO @mduppcl गोंडा जिले के वजीरगंज ब्लॉक में पिछले 12 घंटे से बिजली आपूर्ति पूरी तरह बाधित है।",
            "2026-10-19T06:00:00Z", "19-10-2026 11:30:00", "hi", "HIGH",
            "Concerned team @MVVNLHQ, कृपया सम्मानित उपभोक्ता की बिजली आपूर्ति बाधित होने की शिकायत का संज्ञान लेकर आवश्यक कार्यवाही करें।",
            "DRAFT_READY", "", "", "", "CLASSIFIED", ""
        ],
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/drx_mohd_sami/status/2111617408557568279?s=20",
            "drx_mohd_sami", "upl", "", "YES", "SUPPLY_OUTAGE",
            "PVVNL", "BULANDSAHAR", "BULANDSAHAR", "Verified match for BULANDSAHAR -> PVVNL",
            "Forwarded to Discom", "YES", "", "", "",
            "2111617408557568279",
            "@UPPCLLKO @PVVNLHQ Bulandshahr city power outage since 2 PM. Inverter batteries dying, please restore supply immediately.",
            "2026-10-19T06:15:00Z", "19-10-2026 11:45:00", "en", "HIGH",
            "Concerned team @PVVNLHQ, kindly look into this consumer power supply outage grievance and take prompt action.",
            "DRAFT_READY", "", "", "", "CLASSIFIED", ""
        ],
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/durgeshkumabq8/status/2111609477049012326?s=20",
            "durgeshkumabq8", "akb", "", "YES", "TRANSFORMER_DAMAGE",
            "MVVNL", "AYODHYA", "AYODHYA", "Verified match for AYODHYA -> MVVNL",
            "Forwarded to Discom", "YES", "", "", "",
            "2111609477049012326",
            "@aksharmaBharat @UPPCLLKO हमारे गांव में 63 KVA का ट्रांसफार्मर फुंक गया है। 3 दिन से अंधेरा है। नया ट्रांसफार्मर लगवाया जाए।",
            "2026-10-19T06:20:00Z", "19-10-2026 11:50:00", "hi", "HIGH",
            "Concerned team @MVVNLHQ, कृपया क्षतिग्रस्त परिवर्तक (ट्रांसफार्मर) का अविलंब संज्ञान लेकर नियमानुसार आवश्यक कार्यवाही सुनिश्चित करें।",
            "DRAFT_READY", "", "", "", "CLASSIFIED", ""
        ],
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/Er_Naushad_khan/status/2111626795447706058?s=20",
            "Er_Naushad_khan", "upl", "", "YES", "SUPPLY_VOLTAGE",
            "KESCO", "KESCO", "KANPUR CITY", "Verified match for KANPUR CITY -> KESCO",
            "Forwarded to Discom", "YES", "", "", "",
            "2111626795447706058",
            "@UPPCLLKO @KESCoHQ Rawatpur Kanpur me heavy voltage fluctuation chal raha hai. Fridge and AC damaged. Please fix transformer.",
            "2026-10-19T06:30:00Z", "19-10-2026 12:00:00", "hi-Latn", "HIGH",
            "Concerned team @KESCoHQ, कृपया वोल्टेज समस्या (लो/हाई वोल्टेज) का संज्ञान लेकर आपूर्ति को सामान्य कराना सुनिश्चित करें।",
            "DRAFT_READY", "", "", "", "CLASSIFIED", ""
        ],
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/rohi_tkushwahaa/status/2111631974771360067?s=20",
            "rohi_tkushwahaa", "upm", "", "YES", "CONNECTION_NEW",
            "PuVVNL", "GORAKHPUR-2", "DEORIA", "Verified match for DEORIA -> PuVVNL",
            "Forwarded to Discom", "YES", "", "", "",
            "2111631974771360067",
            "@mduppcl @PuVVNLHQ देवरिया में झटपट पोर्टल पर नए कनेक्शन के लिए 15 दिन पहले अप्लाई किया था लेकिन अभी तक मीटर नहीं लगा।",
            "2026-10-19T06:40:00Z", "19-10-2026 12:10:00", "hi", "HIGH",
            "Concerned team @PuVVNLHQ, कृपया नवीन संयोजन (नए कनेक्शन) सम्बन्धी इस प्रकरण का संज्ञान लेकर आवश्यक कार्यवाही करें।",
            "DRAFT_READY", "", "", "", "CLASSIFIED", ""
        ],
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/PooranSardar/status/2111658518130642993?s=20",
            "PooranSardar", "upl", "", "YES", "BILL_RELATED",
            "DVVNL", "MATHURA", "MATHURA", "Verified match for MATHURA -> DVVNL",
            "Forwarded to Discom", "YES", "", "", "",
            "2111658518130642993",
            "@UPPCLLKO @dvvnlhq मथुरा छाता क्षेत्र में अत्यधिक गलत बिल भेज दिया गया है। 2 कमरे के घर का बिल 45000 आया है।",
            "2026-10-19T06:45:00Z", "19-10-2026 12:15:00", "hi", "HIGH",
            "Concerned team @dvvnlhq, कृपया सम्मानित उपभोक्ता की बिल सम्बन्धी शिकायत का संज्ञान लेकर त्वरित निस्तारण सुनिश्चित करें।",
            "DRAFT_READY", "", "", "", "CLASSIFIED", ""
        ],
        [
            "19-10-2026", "MORNING", "MORNING", "1",
            "https://x.com/Neeraj120459/status/2111660825899323771?s=20",
            "Neeraj120459", "upm", "", "APPRECIATION", "",
            "", "", "", "Consumer appreciation logged",
            "", "", "", "", "",
            "2111660825899323771",
            "@UPPCLLKO @mduppcl Gomti Nagar sub station team resolved my fuse blown complaint within 30 minutes! Excellent work and thank you!",
            "2026-10-19T06:50:00Z", "19-10-2026 12:20:00", "en", "HIGH",
            "", "SKIPPED", "", "", "", "CLASSIFIED", ""
        ]
    ]

    for row_idx, r_data in enumerate(sample_rows, start=3):
        ws_main.row_dimensions[row_idx].height = 45
        for col_idx, val in enumerate(r_data, start=1):
            cell = ws_main.cell(row=row_idx, column=col_idx, value=val)
            cell.border = black_border
            if col_idx == 5:  # LINK URL
                cell.font = link_font
                cell.alignment = left_align
            elif col_idx in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 15, 16, 17, 18, 19, 24, 25, 27, 31):
                cell.font = data_font
                cell.alignment = center_align
            else:
                cell.font = data_font
                cell.alignment = left_align

    # Column dimensions matching original sheet
    column_widths = {
        "A": 15,  # DATE
        "B": 14,  # Link Shift
        "C": 14,  # Reply Shift
        "D": 8,   # C F
        "E": 45,  # LINK URL
        "F": 20,  # Twitter Id
        "G": 10,  # Head
        "H": 20,  # Untagged Complaint
        "I": 15,  # Picked Status
        "J": 26,  # Complaint Type
        "K": 14,  # DISCOM
        "L": 20,  # Zone
        "M": 20,  # District
        "N": 30,  # Remark
        "O": 24,  # Status
        "P": 18,  # REPLY_WITHIN_30_MIN_BY_Laqshya Media
        "Q": 18,  # REPLY_WITHIN_60_MIN_BY_DISCOM
        "R": 18,  # COMPLAINT_NO
        "S": 12,  # STS
        "T": 25,  # Post ID
        "U": 45,  # Post Text
        "V": 24,  # Post Created At (UTC)
        "W": 22,  # First Detected At (IST)
        "X": 12,  # Language
        "Y": 18,  # Routing Confidence
        "Z": 45,  # Suggested Reply Draft
        "AA": 22, # Reply Approval Status
        "AB": 22, # Reply Post ID
        "AC": 35, # Reply Post URL
        "AD": 24, # Reply Posted At
        "AE": 20, # Processing Status
        "AF": 25, # Last Error
    }

    for col_letter, width in column_widths.items():
        ws_main.column_dimensions[col_letter].width = width

    # -------------------------------------------------------------------------
    # TAB 2: Discom_District_List (Authoritative master reference)
    # -------------------------------------------------------------------------
    ws_districts = wb.create_sheet("Discom_District_List")
    ws_districts.append(["DISCOM", "District", "ZONE"])
    ws_districts.row_dimensions[1].height = 30
    for cell in ws_districts[1]:
        cell.font = header_font
        cell.fill = cyan_fill
        cell.alignment = center_align
        cell.border = black_border

    for d_name, d_info in sorted(config.districts.items(), key=lambda x: (x[1].discom, x[0])):
        row_cells = ws_districts.append([d_info.discom, d_name, d_info.zone])
        for c in ws_districts[ws_districts.max_row]:
            c.font = data_font
            c.border = black_border
    
    ws_districts.column_dimensions["A"].width = 16
    ws_districts.column_dimensions["B"].width = 25
    ws_districts.column_dimensions["C"].width = 25

    # -------------------------------------------------------------------------
    # TAB 3: Settings (Operational feature switches & caps)
    # -------------------------------------------------------------------------
    ws_settings = wb.create_sheet("Settings")
    settings_data = [
        ["SETTING_KEY", "VALUE", "DESCRIPTION"],
        ["EMERGENCY_STOP", "FALSE", "Global kill-switch. When TRUE, no polling or replies run."],
        ["AUTO_POST_ENABLED", "FALSE", "When FALSE, replies only draft and require manual approval."],
        ["DAILY_POST_CAP", "50", "Maximum automated replies permitted per calendar day."],
        ["POLLING_INTERVAL_MINS", "5", "Frequency of background post monitoring cycles."],
        ["ALLOWED_DISCOMS", "@PVVNLHQ, @dvvnlhq, @PuVVNLHQ, @KESCoHQ, @MVVNLHQ", "Safety allowlist of receiving handles."],
        ["MONITORED_APEX", "@UPPCLLKO, @EMofficeUP, @AKSharmaOffice, @aksharmaBharat, @ChairmanUppcl, @mduppcl", "Monitored accounts."]
    ]
    for row in settings_data:
        ws_settings.append(row)
    for cell in ws_settings[1]:
        cell.font = header_font
        cell.fill = cyan_fill
        cell.border = black_border
    for row in ws_settings.iter_rows(min_row=2, max_col=3):
        for cell in row:
            cell.font = data_font
            cell.border = black_border
    ws_settings.column_dimensions["A"].width = 25
    ws_settings.column_dimensions["B"].width = 30
    ws_settings.column_dimensions["C"].width = 50

    # -------------------------------------------------------------------------
    # TAB 4: Dashboard
    # -------------------------------------------------------------------------
    ws_dash = wb.create_sheet("Dashboard")
    dash_rows = [
        ["UPPCL X COMPLAINT MONITORING DASHBOARD (TRIAL)", "", ""],
        ["Metric", "Value", "Source / Category"],
        ["Total Complaints Recorded", "=COUNTA('Twitter Complaint List'!E3:E100)", "Active post links count"],
        ["Supply Outages", "=COUNTIF('Twitter Complaint List'!J3:J100, \"SUPPLY_OUTAGE\")", "SUPPLY_OUTAGE"],
        ["Transformer Damage", "=COUNTIF('Twitter Complaint List'!J3:J100, \"TRANSFORMER_DAMAGE\")", "TRANSFORMER_DAMAGE"],
        ["Billing Issues", "=COUNTIF('Twitter Complaint List'!J3:J100, \"BILL_RELATED\")", "BILL_RELATED"],
        ["Voltage Fluctuations", "=COUNTIF('Twitter Complaint List'!J3:J100, \"SUPPLY_VOLTAGE\")", "SUPPLY_VOLTAGE"],
        ["New Connection Issues", "=COUNTIF('Twitter Complaint List'!J3:J100, \"CONNECTION_NEW\")", "CONNECTION_NEW"],
        ["MVVNL Complaints", "=COUNTIF('Twitter Complaint List'!K3:K100, \"MVVNL\")", "Central UP"],
        ["PVVNL Complaints", "=COUNTIF('Twitter Complaint List'!K3:K100, \"PVVNL\")", "Western UP"],
        ["PuVVNL Complaints", "=COUNTIF('Twitter Complaint List'!K3:K100, \"PuVVNL\")", "Eastern UP"],
        ["DVVNL Complaints", "=COUNTIF('Twitter Complaint List'!K3:K100, \"DVVNL\")", "Southern UP"],
        ["KESCo Complaints", "=COUNTIF('Twitter Complaint List'!K3:K100, \"KESCO\")", "Kanpur City"]
    ]
    for row in dash_rows:
        ws_dash.append(row)
    ws_dash.merge_cells("A1:C1")
    ws_dash["A1"].font = Font(name="Verdana", size=14, bold=True, color="1A2F40")
    ws_dash["A1"].fill = cyan_fill
    ws_dash["A1"].alignment = center_align

    for cell in ws_dash[2]:
        cell.font = bold_data_font
        cell.fill = cyan_fill
        cell.border = black_border

    for row in ws_dash.iter_rows(min_row=3, max_col=3):
        for cell in row:
            cell.font = data_font
            cell.border = black_border

    ws_dash.column_dimensions["A"].width = 30
    ws_dash.column_dimensions["B"].width = 20
    ws_dash.column_dimensions["C"].width = 25

    wb.save(output_path)
    print(f"[+] Exact-format workbook generated successfully at: {output_path}")


if __name__ == "__main__":
    build_exact_format_workbook()
